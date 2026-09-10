"""
Document processing module.

Combines:
    - document loading
    - text cleaning
    - text chunking

This module is purely responsible for structural document processing.
It does not handle document authorization/security metadata.
Security metadata is attached during ingestion.
"""

import hashlib
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

ALLOWED_DOCUMENT_EXTENSIONS = {"pdf", "txt", "docx"}

MAX_BLANK_LINES = 2


# ---------------------------------------------------------------------------
# Exceptions
# ---------------------------------------------------------------------------

class DocumentLoadError(Exception):
    """Raised for missing, corrupted, or unreadable documents."""


class UnsupportedFileTypeError(Exception):
    """Raised when the file extension is not supported."""


class EmptyDocumentError(Exception):
    """Raised when a document contains no extractable text."""


# ---------------------------------------------------------------------------
# Loaded document structures
# ---------------------------------------------------------------------------

@dataclass
class LoadedPage:
    page_number: Optional[int]
    text: str


@dataclass
class LoadedDocument:
    filename: str
    file_type: str
    source_path: str
    pages: List[LoadedPage] = field(default_factory=list)

    @property
    def page_count(self) -> Optional[int]:
        return len(self.pages) if self.pages else None


# ---------------------------------------------------------------------------
# Text cleaning
# ---------------------------------------------------------------------------

def clean_text(text: str) -> str:
    """
    Normalize whitespace without changing the actual wording.

    Rules:
        - Normalize CRLF/CR to LF.
        - Remove trailing whitespace from every line.
        - Collapse repeated spaces/tabs.
        - Preserve paragraph breaks.
        - Reduce runs of blank lines to at most MAX_BLANK_LINES.
        - Never summarize, rewrite, or intentionally remove content.
    """

    if not text:
        return ""

    normalized = text.replace("\r\n", "\n").replace("\r", "\n")

    lines = [line.rstrip() for line in normalized.split("\n")]

    cleaned_lines: List[str] = []
    blank_run = 0

    for line in lines:
        if line.strip() == "":
            blank_run += 1

            if blank_run <= MAX_BLANK_LINES:
                cleaned_lines.append("")

        else:
            blank_run = 0

            # Correct regex:
            # multiple spaces/tabs -> one space
            cleaned_lines.append(
                re.sub(r"[ \t]+", " ", line).strip()
            )

    return "\n".join(cleaned_lines).strip()


# ---------------------------------------------------------------------------
# Document loader
# ---------------------------------------------------------------------------

class DocumentLoader:
    """
    Loads PDF, DOCX and TXT documents.

    OCR is intentionally not included in this phase.
    Scanned PDFs without an extractable text layer will produce
    empty page text and ultimately raise EmptyDocumentError.
    """

    def load(
        self,
        file_path: str,
        file_type: Optional[str] = None,
    ) -> LoadedDocument:

        if not os.path.isfile(file_path):
            raise DocumentLoadError(
                f"File not found: {file_path}"
            )

        extension = (
            file_type
            or os.path.splitext(file_path)[1].lstrip(".")
        ).lower()

        if extension not in ALLOWED_DOCUMENT_EXTENSIONS:
            raise UnsupportedFileTypeError(
                f"Unsupported file type '.{extension}'. "
                f"Supported types: "
                f"{', '.join(sorted(ALLOWED_DOCUMENT_EXTENSIONS))}."
            )

        filename = os.path.basename(file_path)

        if extension == "pdf":
            pages = self._load_pdf(file_path)

        elif extension == "docx":
            pages = self._load_docx(file_path)

        elif extension == "txt":
            pages = self._load_txt(file_path)

        else:
            raise UnsupportedFileTypeError(
                f"No loader implemented for '.{extension}'"
            )

        if not any(page.text.strip() for page in pages):
            raise EmptyDocumentError(
                f"No extractable text was found in '{filename}'. "
                "If this is a scanned/image-based document, OCR "
                "support is not available in this phase."
            )

        return LoadedDocument(
            filename=filename,
            file_type=extension,
            source_path=file_path,
            pages=pages,
        )

    # -----------------------------------------------------------------------
    # PDF
    # -----------------------------------------------------------------------

    def _load_pdf(self, file_path: str) -> List[LoadedPage]:

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentLoadError(
                "pypdf is not installed. "
                "Run 'pip install -r requirements.txt'."
            ) from exc

        try:
            reader = PdfReader(file_path)
        except Exception as exc:
            raise DocumentLoadError(
                f"Failed to open PDF '{file_path}': {exc}"
            ) from exc

        if reader.is_encrypted:
            raise DocumentLoadError(
                f"PDF '{file_path}' is password-protected "
                "and cannot be indexed."
            )

        pages: List[LoadedPage] = []

        for index, page in enumerate(reader.pages, start=1):

            try:
                text = page.extract_text() or ""

            except Exception as exc:
                raise DocumentLoadError(
                    f"Failed to extract text from page {index}: {exc}"
                ) from exc

            pages.append(
                LoadedPage(
                    page_number=index,
                    text=text,
                )
            )

        return pages

    # -----------------------------------------------------------------------
    # DOCX
    # -----------------------------------------------------------------------

    def _load_docx(self, file_path: str) -> List[LoadedPage]:

        try:
            from docx import Document as DocxDocument

        except ImportError as exc:
            raise DocumentLoadError(
                "python-docx is not installed. "
                "Run 'pip install -r requirements.txt'."
            ) from exc

        try:
            docx_document = DocxDocument(file_path)

        except Exception as exc:
            raise DocumentLoadError(
                f"Failed to open DOCX '{file_path}': {exc}"
            ) from exc

        parts: List[str] = []

        # Paragraphs
        for paragraph in docx_document.paragraphs:

            if paragraph.text.strip():
                parts.append(paragraph.text)

        # Tables
        for table in docx_document.tables:

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(" | ".join(cells))

        # DOCX does not provide reliable page boundaries
        # without a rendering engine.
        return [
            LoadedPage(
                page_number=None,
                text="\n".join(parts),
            )
        ]

    # -----------------------------------------------------------------------
    # TXT
    # -----------------------------------------------------------------------

    def _load_txt(self, file_path: str) -> List[LoadedPage]:

        for encoding in ("utf-8", "latin-1"):

            try:

                with open(
                    file_path,
                    "r",
                    encoding=encoding,
                ) as handle:

                    text = handle.read()

                return [
                    LoadedPage(
                        page_number=None,
                        text=text,
                    )
                ]

            except UnicodeDecodeError:
                continue

            except OSError as exc:
                raise DocumentLoadError(
                    f"Failed to read TXT file '{file_path}': {exc}"
                ) from exc

        raise DocumentLoadError(
            "Could not decode TXT file with supported encodings."
        )


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

@dataclass
class TextChunk:
    """
    Represents one structurally generated text chunk.
    """

    page_number: Optional[int]
    chunk_index: int
    text: str


class Chunker:
    """
    Splits cleaned page text into overlapping chunks.

    Chunking is based on words rather than characters so that chunks
    remain word-boundary-safe.

    Environment variables:

        RAG_CHUNK_SIZE
            Number of words per chunk.

        RAG_CHUNK_OVERLAP
            Number of overlapping words between chunks.
    """

    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
    ):

        self.chunk_size = chunk_size or int(
            os.getenv("RAG_CHUNK_SIZE", "500")
        )

        self.chunk_overlap = chunk_overlap or int(
            os.getenv("RAG_CHUNK_OVERLAP", "50")
        )

        if self.chunk_size <= 0:
            raise ValueError(
                "RAG_CHUNK_SIZE must be greater than zero."
            )

        if self.chunk_overlap < 0:
            raise ValueError(
                "RAG_CHUNK_OVERLAP cannot be negative."
            )

        if self.chunk_overlap >= self.chunk_size:
            raise ValueError(
                "RAG_CHUNK_OVERLAP must be smaller than RAG_CHUNK_SIZE."
            )

    def chunk_text(
        self,
        text: str,
        page_number: Optional[int] = None,
    ) -> List[TextChunk]:
        """
        Chunk a single page of text.
        """

        cleaned = clean_text(text)

        if not cleaned:
            return []

        words = cleaned.split()

        if not words:
            return []

        chunks: List[TextChunk] = []

        step = self.chunk_size - self.chunk_overlap

        chunk_index = 0

        for start in range(0, len(words), step):

            end = min(
                start + self.chunk_size,
                len(words),
            )

            chunk_words = words[start:end]

            if not chunk_words:
                continue

            chunk_text = " ".join(chunk_words)

            chunks.append(
                TextChunk(
                    page_number=page_number,
                    chunk_index=chunk_index,
                    text=chunk_text,
                )
            )

            chunk_index += 1

            if end >= len(words):
                break

        return chunks

    def chunk_pages(
        self,
        pages: List[Tuple[Optional[int], str]],
    ) -> List[TextChunk]:
        """
        Chunk multiple pages independently.

        Chunk indexes restart from zero for each page.
        """

        all_chunks: List[TextChunk] = []

        for page_number, text in pages:

            page_chunks = self.chunk_text(
                text=text,
                page_number=page_number,
            )

            all_chunks.extend(page_chunks)

        return all_chunks


# ---------------------------------------------------------------------------
# Stable chunk ID
# ---------------------------------------------------------------------------

def generate_chunk_id(
    document_checksum: str,
    page_number: Optional[int],
    chunk_index: int,
) -> str:
    """
    Generate a deterministic chunk ID.

    Same document checksum + page + chunk index
    always produces the same chunk ID.
    """

    raw = (
        f"{document_checksum}:"
        f"{page_number}:"
        f"{chunk_index}"
    )

    return hashlib.sha256(
        raw.encode("utf-8")
    ).hexdigest()